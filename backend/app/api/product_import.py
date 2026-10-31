import csv
import io
import re
import uuid
from collections import defaultdict
from dataclasses import dataclass
from itertools import islice
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.concurrency import run_in_threadpool

from app.api.deps import get_current_outlet, require_roles
from app.core.constants import MAX_BIGINT
from app.db.session import get_db
from app.models.category import Category
from app.models.product import Product
from app.models.stock_movement import StockMovement
from app.models.user import User
from app.models.outlet import Outlet
from app.schemas.product_import import ImportResult, RowMessage

router = APIRouter(prefix="/products/import", tags=["products-import"])

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
MAX_INT = MAX_BIGINT
MAX_ERRORS_SHOWN = 200
MAX_SHEETS = 20

HEADER_ALIASES = {
    "name": "name",
    "nama": "name",
    "nama produk": "name",
    "price": "price",
    "harga": "price",
    "harga jual": "price",
    "sku": "sku",
    "barcode": "barcode",
    "stock": "stock",
    "stok": "stock",
    "category": "category",
    "kategori": "category",
}
REQUIRED_COLUMNS = {"name", "price"}

TEMPLATE_CSV = (
    "name,price,sku,barcode,stock,category\r\n"
    "Teh Botol,5000,TB-001,8991234567890,50,Minuman\r\n"
    "Roti Tawar,15000,RT-001,,20,Makanan\r\n"
)


@dataclass
class ParsedRow:
    row: int
    name: str
    price: int
    sku: str | None
    barcode: str | None
    stock: int | None
    category: str | None
    sheet: str | None = None


@dataclass
class PlanItem:
    row: ParsedRow
    target: Product | None


def cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


NUMERIC_IDENTIFIER_MARKER = "\x00NUMERIC_IDENTIFIER\x00"


def excel_identifier_to_text(value: Any, number_format: str | None) -> str:
    """Keep zero-padding from a pure Excel zero mask; flag ambiguous numeric IDs."""
    if value is None:
        return ""
    if number_format and re.fullmatch(r"0+", number_format):
        try:
            number = int(value)
            if number == value and 0 <= number < 10**15 and len(number_format) <= 15:
                return str(number).zfill(len(number_format))
        except (TypeError, ValueError, OverflowError):
            pass
    return NUMERIC_IDENTIFIER_MARKER + cell_to_text(value)


def identifier_columns(header: list[str]) -> set[int]:
    return {
        index
        for index, value in enumerate(header)
        if HEADER_ALIASES.get(value.strip().lower()) in {"sku", "barcode"}
    }


def read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")
    lines = text.splitlines()
    first_line = lines[0] if lines else ""
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [[cell.strip() for cell in row] for row in reader]


def read_excel_sheets(content: bytes, extension: str) -> list[tuple[str, list[list[str]]]]:
    try:
        if extension == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
            try:
                if len(workbook.worksheets) > MAX_SHEETS:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Maksimal {MAX_SHEETS} sheet per file Excel",
                    )
                sheets = []
                for sheet in workbook.worksheets:
                    rows = []
                    identifiers: set[int] = set()
                    for row_index, row in enumerate(
                        islice(sheet.iter_rows(), MAX_ROWS + 1000)
                    ):
                        if row_index == 0:
                            identifiers = identifier_columns(
                                [cell_to_text(cell.value) for cell in row]
                            )
                        rows.append([
                            excel_identifier_to_text(cell.value, cell.number_format)
                            if index in identifiers and cell.data_type == "n"
                            else cell_to_text(cell.value)
                            for index, cell in enumerate(row)
                        ])
                    sheets.append((sheet.title, rows))
                return sheets
            finally:
                workbook.close()

        import xlrd

        workbook = xlrd.open_workbook(
            file_contents=content, on_demand=True, formatting_info=True
        )
        try:
            if workbook.nsheets > MAX_SHEETS:
                raise HTTPException(
                    status_code=400,
                    detail=f"Maksimal {MAX_SHEETS} sheet per file Excel",
                )
            sheets = []
            for sheet_index in range(workbook.nsheets):
                sheet = workbook.sheet_by_index(sheet_index)
                rows = []
                identifiers: set[int] = set()
                for row_index in range(min(sheet.nrows, MAX_ROWS + 1000)):
                    cells = sheet.row(row_index)
                    if row_index == 0:
                        identifiers = identifier_columns(
                            [cell_to_text(cell.value) for cell in cells]
                        )
                    converted = []
                    for column_index, cell in enumerate(cells):
                        if column_index in identifiers and cell.ctype == xlrd.XL_CELL_NUMBER:
                            xf = workbook.xf_list[cell.xf_index]
                            number_format = workbook.format_map[xf.format_key].format_str
                            converted.append(excel_identifier_to_text(cell.value, number_format))
                        else:
                            converted.append(cell_to_text(cell.value))
                    rows.append(converted)
                sheets.append((sheet.name, rows))
            return sheets
        finally:
            workbook.release_resources()
    except HTTPException:
        raise
    except ImportError:
        raise HTTPException(
            status_code=500, detail="Dukungan file .xls belum terpasang di server"
        )
    except Exception:
        raise HTTPException(status_code=400, detail="File Excel tidak dapat dibaca")


def to_records(rows: list[list[str]]) -> list[dict[str, Any]]:
    if not rows:
        raise HTTPException(status_code=400, detail="File kosong")

    columns: dict[int, str] = {}
    for index, header in enumerate(rows[0]):
        key = HEADER_ALIASES.get(header.strip().lower())
        if key and key not in columns.values():
            columns[index] = key

    missing = REQUIRED_COLUMNS - set(columns.values())
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Kolom wajib tidak ditemukan: {', '.join(sorted(missing))}",
        )

    records: list[dict[str, Any]] = []
    for line_no, row in enumerate(rows[1:], start=2):
        if not any(cell for cell in row):
            continue
        record: dict[str, Any] = {"row": line_no}
        for index, key in columns.items():
            record[key] = row[index] if index < len(row) else ""
        records.append(record)

    if not records:
        raise HTTPException(status_code=400, detail="Tidak ada baris data")
    if len(records) > MAX_ROWS:
        raise HTTPException(
            status_code=400, detail=f"Maksimal {MAX_ROWS} baris data per file"
        )
    return records


def to_records_from_sheets(
    sheets: list[tuple[str, list[list[str]]]],
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    row_offset = 0
    for sheet_name, rows in sheets:
        if not rows or not any(cell for row in rows[1:] for cell in row):
            continue
        try:
            sheet_records = to_records(rows)
        except HTTPException as exc:
            raise HTTPException(
                status_code=exc.status_code,
                detail=f"Sheet {sheet_name}: {exc.detail}",
            )
        for record in sheet_records:
            record["row"] += row_offset
            record["sheet"] = sheet_name
        records.extend(sheet_records)
        row_offset += len(sheet_records)
        if len(records) > MAX_ROWS:
            raise HTTPException(
                status_code=400, detail=f"Maksimal {MAX_ROWS} baris data per file"
            )
    if not records:
        raise HTTPException(status_code=400, detail="Tidak ada baris data di semua sheet")
    return records


def parse_int(text: str, label: str) -> tuple[int | None, str | None]:
    cleaned = text.strip()
    if cleaned.lower().startswith("rp"):
        cleaned = cleaned[2:].strip()
    if not (cleaned.isascii() and cleaned.isdigit()):
        return None, f"{label} harus angka bulat (tanpa titik, koma, atau huruf)"
    value = int(cleaned)
    if value > MAX_INT:
        return None, f"{label} terlalu besar"
    return value, None


def validate_records(
    records: list[dict[str, Any]],
) -> tuple[list[ParsedRow], list[RowMessage]]:
    parsed: list[ParsedRow] = []
    errors: list[RowMessage] = []
    seen_sku: dict[str, int] = {}
    seen_barcode: dict[str, int] = {}
    seen_name: dict[str, int] = {}

    for rec in records:
        row = rec["row"]
        problems: list[str] = []

        name = rec.get("name", "").strip()
        if not name:
            problems.append("nama produk kosong")
        elif len(name) > 150:
            problems.append("nama produk maksimal 150 karakter")

        price, price_error = parse_int(rec.get("price", ""), "harga")
        if price_error:
            problems.append(price_error)

        sku_text = rec.get("sku", "").strip()
        if sku_text.startswith(NUMERIC_IDENTIFIER_MARKER):
            problems.append("SKU numerik harus disimpan sebagai teks di Excel")
        sku = sku_text or None
        if sku and len(sku) > 50:
            problems.append("SKU maksimal 50 karakter")

        barcode_text = rec.get("barcode", "").strip()
        if barcode_text.startswith(NUMERIC_IDENTIFIER_MARKER):
            problems.append("barcode numerik harus disimpan sebagai teks di Excel")
        barcode = barcode_text or None
        if barcode and len(barcode) > 50:
            problems.append("barcode maksimal 50 karakter")

        stock: int | None = None
        stock_text = rec.get("stock", "").strip()
        if stock_text:
            stock, stock_error = parse_int(stock_text, "stok")
            if stock_error:
                problems.append(stock_error)

        category = rec.get("category", "").strip() or None
        if category and len(category) > 100:
            problems.append("nama kategori maksimal 100 karakter")

        if sku and sku in seen_sku:
            problems.append(f"SKU dobel dengan baris {seen_sku[sku]}")
        if barcode and barcode in seen_barcode:
            problems.append(f"barcode dobel dengan baris {seen_barcode[barcode]}")
        if not sku and not barcode and name and name.lower() in seen_name:
            problems.append(
                f"nama dobel dengan baris {seen_name[name.lower()]} "
                "(tambahkan SKU atau barcode)"
            )

        if problems:
            sheet = rec.get("sheet")
            prefix = f"Sheet {sheet}: " if sheet else ""
            errors.append(RowMessage(row=row, message=prefix + "; ".join(problems)))
            continue

        if sku:
            seen_sku[sku] = row
        if barcode:
            seen_barcode[barcode] = row
        if not sku and not barcode:
            seen_name[name.lower()] = row

        parsed.append(
            ParsedRow(
                row=row,
                name=name,
                price=price,
                sku=sku,
                barcode=barcode,
                stock=stock,
                category=category,
                sheet=rec.get("sheet"),
            )
        )

    return parsed, errors


async def build_plan(
    db: AsyncSession, parsed: list[ParsedRow], errors: list[RowMessage], outlet_id: uuid.UUID
) -> list[PlanItem]:
    skus = {r.sku for r in parsed if r.sku}
    barcodes = {r.barcode for r in parsed if r.barcode}
    names = {r.name.lower() for r in parsed if not r.sku and not r.barcode}

    conditions = []
    if skus:
        conditions.append(Product.sku.in_(skus))
    if barcodes:
        conditions.append(Product.barcode.in_(barcodes))
    if names:
        conditions.append(func.lower(Product.name).in_(names))

    existing: list[Product] = []
    if conditions:
        result = await db.execute(
            select(Product).where(Product.outlet_id == outlet_id, Product.deleted_at.is_(None), or_(*conditions))
        )
        existing = list(result.scalars().all())

    sku_map = {p.sku: p for p in existing if p.sku}
    barcode_map = {p.barcode: p for p in existing if p.barcode}
    name_map: dict[str, list[Product]] = defaultdict(list)
    for product in existing:
        name_map[product.name.lower()].append(product)

    plan: list[PlanItem] = []
    used_targets: dict[uuid.UUID, int] = {}

    for r in parsed:
        by_sku = sku_map.get(r.sku) if r.sku else None
        by_barcode = barcode_map.get(r.barcode) if r.barcode else None

        if by_sku and by_barcode and by_sku.id != by_barcode.id:
            errors.append(
                RowMessage(
                    row=r.row,
                    message="SKU dan barcode milik dua produk yang berbeda",
                )
            )
            continue

        target = by_sku or by_barcode
        if target is None and not r.sku and not r.barcode:
            matches = name_map.get(r.name.lower(), [])
            if len(matches) > 1:
                errors.append(
                    RowMessage(
                        row=r.row,
                        message=(
                            "ada lebih dari satu produk dengan nama ini, "
                            "tambahkan SKU atau barcode"
                        ),
                    )
                )
                continue
            target = matches[0] if matches else None

        if target is not None:
            if target.id in used_targets:
                errors.append(
                    RowMessage(
                        row=r.row,
                        message=(
                            "produk yang sama sudah diproses di baris "
                            f"{used_targets[target.id]}"
                        ),
                    )
                )
                continue
            used_targets[target.id] = r.row

        plan.append(PlanItem(row=r, target=target))

    return plan


async def resolve_categories(
    db: AsyncSession, parsed: list[ParsedRow], outlet_id: uuid.UUID
) -> tuple[dict[str, uuid.UUID], list[str]]:
    wanted: dict[str, str] = {}
    for r in parsed:
        if r.category and r.category.lower() not in wanted:
            wanted[r.category.lower()] = r.category
    if not wanted:
        return {}, []

    result = await db.execute(
        select(Category).where(
            Category.deleted_at.is_(None),
            Category.outlet_id == outlet_id,
            func.lower(Category.name).in_(list(wanted.keys())),
        )
    )
    existing = {c.name.lower(): c.id for c in result.scalars().all()}
    missing = [name for key, name in wanted.items() if key not in existing]
    return existing, missing


async def apply_plan(
    db: AsyncSession,
    plan: list[PlanItem],
    category_ids: dict[str, uuid.UUID],
    new_categories: list[str],
    user: User,
    outlet: Outlet,
) -> None:
    try:
        for name in new_categories:
            category = Category(id=uuid.uuid4(), name=name, outlet_id=outlet.id)
            db.add(category)
            category_ids[name.lower()] = category.id
        await db.flush()

        movements: list[StockMovement] = []
        for item in plan:
            r = item.row
            category_id = category_ids.get(r.category.lower()) if r.category else None

            if item.target is None:
                product = Product(
                    id=uuid.uuid4(),
                    outlet_id=outlet.id,
                    name=r.name,
                    sku=r.sku,
                    barcode=r.barcode,
                    price=r.price,
                    stock=r.stock or 0,
                    category_id=category_id,
                )
                db.add(product)
                if r.stock:
                    movements.append(
                        StockMovement(
                            product_id=product.id,
                            outlet_id=outlet.id,
                            change=r.stock,
                            stock_after=r.stock,
                            reason="initial",
                            note="Impor file",
                            user_id=user.id,
                        )
                    )
            else:
                target = item.target
                target.name = r.name
                target.price = r.price
                if r.sku:
                    target.sku = r.sku
                if r.barcode:
                    target.barcode = r.barcode
                if category_id is not None:
                    target.category_id = category_id

        await db.flush()
        db.add_all(movements)
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Impor gagal karena ada data yang bentrok. Jalankan pratinjau lagi.",
        )


@router.get("/template")
async def download_template(
    _: User = Depends(require_roles("owner", "admin")),
    __: Outlet = Depends(get_current_outlet),
):
    content = b"\xef\xbb\xbf" + TEMPLATE_CSV.encode("utf-8")
    return Response(
        content=content,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="template_produk.csv"'},
    )


@router.post("", response_model=ImportResult)
async def import_products(
    file: UploadFile = File(...),
    dry_run: bool = Query(
        default=True,
        description="true: hanya periksa. false: simpan jika tidak ada error.",
    ),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_roles("owner", "admin")),
    outlet: Outlet = Depends(get_current_outlet),
):
    filename = (file.filename or "").lower()
    if filename.endswith(".csv"):
        content = await file.read(MAX_BYTES + 1)
        if len(content) > MAX_BYTES:
            raise HTTPException(status_code=413, detail="Ukuran file maksimal 5 MB")
        records = to_records(read_csv(content))
    elif filename.endswith(".xlsx"):
        extension = ".xlsx"
    else:
        if filename.endswith(".xls"):
            extension = ".xls"
        else:
            raise HTTPException(
                status_code=400, detail="Format file harus .csv, .xlsx, atau .xls"
            )
    if not filename.endswith(".csv"):
        content = await file.read(MAX_BYTES + 1)
        if len(content) > MAX_BYTES:
            raise HTTPException(status_code=413, detail="Ukuran file maksimal 5 MB")
        sheets = await run_in_threadpool(read_excel_sheets, content, extension)
        records = to_records_from_sheets(sheets)

    parsed, errors = validate_records(records)
    plan = await build_plan(db, parsed, errors, outlet.id)
    category_ids, new_categories = await resolve_categories(db, parsed, outlet.id)

    errors.sort(key=lambda e: e.row)
    warnings = [
        RowMessage(
            row=item.row.row,
            message=(
                f"Sheet {item.row.sheet}: " if item.row.sheet else ""
            )
            + "stok diabaikan karena produk sudah ada, gunakan penyesuaian stok",
        )
        for item in plan
        if item.target is not None and item.row.stock
    ]

    to_create = sum(1 for item in plan if item.target is None)
    to_update = len(plan) - to_create

    saved = False
    if not dry_run and not errors:
        await apply_plan(db, plan, category_ids, new_categories, user, outlet)
        saved = True

    return ImportResult(
        dry_run=dry_run,
        saved=saved,
        total_rows=len(records),
        to_create=to_create,
        to_update=to_update,
        categories_to_create=new_categories,
        error_count=len(errors),
        errors=errors[:MAX_ERRORS_SHOWN],
        warnings=warnings[:MAX_ERRORS_SHOWN],
    )
